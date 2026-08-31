from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Guia para instalar Git y ejecutar el proyecto.docx"


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "Courier New"
        r.font.size = Pt(9)
        if i < len(lines) - 1:
            r.add_break()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, before, after in [
        ("Heading 1", 20, 16, 6),
        ("Heading 2", 16, 12, 4),
        ("Heading 3", 14, 8, 3),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Guía Para Instalar Git Y Ejecutar El Proyecto")
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph("Proyecto: Sistema Integrador 6to PPP")

    heading(doc, "Objetivo", 1)
    doc.add_paragraph(
        "Esta guía explica cómo preparar la computadora para descargar el proyecto desde GitHub y ejecutar el frontend."
    )

    heading(doc, "1. Instalar Git", 1)
    numbered(
        doc,
        [
            "Entrar a https://git-scm.com/downloads",
            "Elegir Windows.",
            "Descargar el instalador.",
            "Ejecutar el instalador y dejar las opciones por defecto.",
            "Finalizar la instalación.",
        ],
    )

    heading(doc, "2. Abrir Git Bash", 1)
    doc.add_paragraph("Git Bash es la terminal que se instala junto con Git en Windows.")
    bullets(
        doc,
        [
            "Opción 1: abrir el menú inicio, escribir Git Bash y hacer clic en la aplicación.",
            "Opción 2: entrar a una carpeta, hacer clic derecho y elegir Open Git Bash here.",
            "Opción 3: desde Visual Studio Code, abrir Terminal > New Terminal y elegir Git Bash.",
        ],
    )

    heading(doc, "3. Verificar Que Git Funcione", 1)
    doc.add_paragraph("En Git Bash escribir:")
    code(doc, ["git --version"])
    doc.add_paragraph("Si aparece una versión de Git, la instalación está correcta.")

    heading(doc, "4. Configurar Nombre Y Email", 1)
    doc.add_paragraph("Esto se hace una sola vez por computadora.")
    code(
        doc,
        [
            'git config --global user.name "Nombre Apellido"',
            'git config --global user.email "correo@email.com"',
        ],
    )
    doc.add_paragraph("Ejemplo:")
    code(
        doc,
        [
            'git config --global user.name "Juan Perez"',
            'git config --global user.email "juanperez@gmail.com"',
        ],
    )

    heading(doc, "5. Instalar Node.js", 1)
    numbered(
        doc,
        [
            "Entrar a https://nodejs.org",
            "Descargar la versión recomendada para la mayoría de los usuarios.",
            "Instalar dejando las opciones por defecto.",
            "Cerrar y volver a abrir Git Bash.",
        ],
    )
    doc.add_paragraph("Verificar la instalación:")
    code(doc, ["node --version", "npm --version"])

    heading(doc, "6. Clonar El Repositorio", 1)
    doc.add_paragraph("Clonar significa descargar una copia del proyecto desde GitHub a la computadora.")
    doc.add_paragraph("Primero ubicarse en una carpeta donde guardar el proyecto. Por ejemplo:")
    code(doc, ["cd Desktop"])
    doc.add_paragraph("Después ejecutar:")
    code(doc, ["git clone https://github.com/Silmorato/sistema-integrador-6to-ppp.git"])

    heading(doc, "7. Abrir El Proyecto", 1)
    code(doc, ["cd sistema-integrador-6to-ppp", "code ."])
    doc.add_paragraph("Si el comando code . no funciona, abrir Visual Studio Code y elegir File > Open Folder.")

    heading(doc, "8. Ejecutar El Frontend", 1)
    doc.add_paragraph("El frontend está dentro de la carpeta frontend.")
    code(doc, ["cd frontend", "npm install", "npm run dev"])
    doc.add_paragraph(
        "Cuando aparezca una dirección parecida a http://localhost:5173, abrirla en el navegador."
    )

    heading(doc, "Resumen De Comandos", 1)
    code(
        doc,
        [
            "git --version",
            'git config --global user.name "Nombre Apellido"',
            'git config --global user.email "correo@email.com"',
            "git clone https://github.com/Silmorato/sistema-integrador-6to-ppp.git",
            "cd sistema-integrador-6to-ppp",
            "cd frontend",
            "npm install",
            "npm run dev",
        ],
    )

    heading(doc, "Ideas Clave", 1)
    bullets(
        doc,
        [
            "Git guarda versiones del proyecto en la computadora.",
            "GitHub permite compartir el proyecto en internet.",
            "Clonar es descargar el proyecto desde GitHub.",
            "npm install instala las herramientas necesarias del frontend.",
            "npm run dev inicia el proyecto para verlo en el navegador.",
        ],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
