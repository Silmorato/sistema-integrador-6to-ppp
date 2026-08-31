from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Proyecto Integrador - Consignas generales para hoy.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_text(cell, text, bold=False, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = "Arial"
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "Courier New"
        r.font.size = Pt(9)
        if i < len(lines) - 1:
            r.add_break()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, before, after in [
        ("Heading 1", 20, 16, 6),
        ("Heading 2", 16, 12, 4),
        ("Heading 3", 14, 8, 3),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Proyecto Integrador: Consignas Generales Para Hoy")
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph("Actividad común para todos los grupos.")

    heading(doc, "Idea De Trabajo", 1)
    doc.add_paragraph(
        "Hoy vamos a organizar cómo se va a construir el frontend del sistema integrador. "
        "La actividad es la misma para todos los grupos: cada grupo debe analizar su módulo y completar la ficha de trabajo."
    )
    code(
        doc,
        [
            "No vamos a construir cuatro sistemas separados.",
            "Vamos a construir un solo sistema integrador dividido en módulos.",
        ],
    )

    heading(doc, "Objetivo De La Clase", 1)
    bullets(
        doc,
        [
            "Identificar qué debe hacer cada módulo.",
            "Definir pantallas, datos, acciones, estados y roles.",
            "Acordar una forma común de trabajo entre todos los grupos.",
            "Preparar el proyecto para que después pueda integrarse.",
        ],
    )

    heading(doc, "Módulos Del Sistema", 1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = table.rows[0].cells
    for idx, h in enumerate(["Grupo", "Módulo", "Responsabilidad principal"]):
        set_text(headers[idx], h, bold=True)
        shade(headers[idx], "F2F4F7")

    rows = [
        ("Grupo 1", "Autenticación y sistema general", "Login, inicio, menú, usuarios, roles y acceso a los módulos."),
        ("Grupo 2", "Laboratorio", "Carga, listado y gestión de turnos de laboratorio."),
        ("Grupo 3", "Comedor", "Carga, listado y gestión de reservas o turnos de comedor."),
        ("Grupo 4", "Notebooks y accesorios", "Registro y consulta de notebooks, accesorios, estados y movimientos de gabinete."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_text(cells[idx], value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    heading(doc, "Consigna General", 1)
    doc.add_paragraph("Cada grupo debe trabajar sobre el módulo que le corresponde y responder:")
    numbered(
        doc,
        [
            "Qué problema resuelve el módulo.",
            "Qué pantallas necesita.",
            "Qué datos se cargan.",
            "Qué datos se muestran.",
            "Qué acciones puede realizar el usuario.",
            "Qué estados puede tener la información cargada.",
            "Qué roles pueden usar el módulo.",
            "Qué información necesita de otros módulos.",
        ],
    )

    doc.add_page_break()
    heading(doc, "Preguntas Orientadoras Por Módulo", 1)

    heading(doc, "Autenticación Y Sistema General", 2)
    bullets(
        doc,
        [
            "¿Cómo ingresa un usuario al sistema?",
            "¿Qué datos pide el login?",
            "¿Qué aparece en la pantalla de inicio?",
            "¿Cómo se accede a cada módulo?",
            "¿Qué roles existen?",
            "¿Qué puede ver o hacer cada rol?",
            "¿Habrá novedades, calendario o accesos rápidos?",
        ],
    )

    heading(doc, "Laboratorio", 2)
    bullets(
        doc,
        [
            "¿Quién puede pedir un turno?",
            "¿Quién puede confirmarlo?",
            "¿Qué datos son obligatorios?",
            "¿Cómo se muestra el listado?",
            "¿Qué pasa si se cancela un turno?",
        ],
    )

    heading(doc, "Comedor", 2)
    bullets(
        doc,
        [
            "¿Quién puede cargar una reserva?",
            "¿Qué información necesita comedor?",
            "¿Cómo se controla la cantidad de alumnos?",
            "¿Qué datos se muestran en el listado?",
            "¿Qué pasa si una reserva se cancela?",
        ],
    )

    heading(doc, "Notebooks Y Accesorios", 2)
    bullets(
        doc,
        [
            "¿Cómo se identifica cada notebook o accesorio?",
            "¿Qué estados puede tener?",
            "¿Quién puede cambiar el estado?",
            "¿A quién puede estar asignada?",
            "¿Cómo se muestra si está disponible, en uso o en reparación?",
            "¿Qué movimientos conviene registrar?",
        ],
    )

    heading(doc, "Acuerdos Para Todos Los Grupos", 1)
    bullets(
        doc,
        [
            "Usar nombres claros en español.",
            "Pensar el sistema como una sola aplicación.",
            "Mantener pantallas parecidas entre los módulos.",
            "Usar los mismos nombres para botones comunes.",
            "Definir roles y permisos visuales.",
            "Avisar si un módulo necesita información de otro.",
        ],
    )
    doc.add_paragraph("Botones comunes sugeridos: Nuevo, Guardar, Cancelar, Editar, Eliminar, Ver detalle y Volver.")

    doc.add_page_break()
    heading(doc, "Ficha Para Completar", 1)
    doc.add_paragraph("Cada grupo debe completar esta ficha durante la clase.")
    code(
        doc,
        [
            "Nombre del módulo:",
            "",
            "Integrantes:",
            "",
            "Objetivo del módulo:",
            "",
            "Problema que resuelve:",
            "",
            "Pantallas necesarias:",
            "",
            "Datos que se cargan:",
            "",
            "Datos que se muestran:",
            "",
            "Acciones del usuario:",
            "",
            "Estados posibles:",
            "",
            "Roles que usan este módulo:",
            "",
            "Información que necesita de otros módulos:",
            "",
            "Dudas o decisiones pendientes:",
        ],
    )

    heading(doc, "Puesta En Común", 1)
    doc.add_paragraph("Al finalizar la actividad, cada grupo comparte brevemente:")
    bullets(
        doc,
        [
            "Qué módulo le tocó.",
            "Qué pantallas pensó.",
            "Qué datos necesita cargar.",
            "Qué acciones puede realizar el usuario.",
            "Qué roles usarían ese módulo.",
            "Qué dudas aparecieron.",
        ],
    )

    heading(doc, "Entrega Para La Próxima Clase", 1)
    bullets(
        doc,
        [
            "Ficha del módulo completa.",
            "Boceto simple de las pantallas principales.",
            "Lista de dudas o decisiones para consultar.",
        ],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
