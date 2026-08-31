from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Proyecto Integrador - Actividad para estudiantes.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(item)


def add_code_block(doc, lines):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        if index < len(lines) - 1:
            run.add_break()


def add_group_card(doc, group, module, responsibility, expected):
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    rows = [
        ("Grupo", group),
        ("Módulo", module),
        ("Responsabilidad", responsibility),
        ("Debe pensar", expected),
    ]
    for row_index, (left, right) in enumerate(rows):
        cells = table.rows[row_index].cells
        set_cell_text(cells[0], left, bold=True)
        set_cell_text(cells[1], right)
        cells[0].width = Inches(1.6)
        cells[1].width = Inches(4.9)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_index == 0:
            set_cell_shading(cells[0], "F2F4F7")
            set_cell_shading(cells[1], "F2F4F7")
    doc.add_paragraph()


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, before, after in [
        ("Heading 1", 20, 18, 6),
        ("Heading 2", 16, 14, 4),
        ("Heading 3", 14, 10, 3),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Proyecto Integrador: Organización Del Trabajo")
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph("Actividad para comenzar a definir los módulos del sistema.")

    add_heading(doc, "Idea Principal", level=1)
    doc.add_paragraph(
        "Vamos a construir el frontend de un sistema integrador para gestionar distintos movimientos de la escuela."
    )
    add_code_block(
        doc,
        [
            "No vamos a construir cuatro sistemas separados.",
            "Vamos a construir un solo sistema dividido en módulos.",
        ],
    )

    add_heading(doc, "Qué Vamos A Hacer", level=1)
    add_bullets(
        doc,
        [
            "Pensar qué debe hacer cada módulo.",
            "Definir pantallas, datos, acciones y estados.",
            "Acordar una forma común de trabajar.",
            "Preparar el proyecto para que después pueda integrarse.",
            "Trabajar primero con datos simulados.",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "Grupos Y Módulos", level=1)
    add_group_card(
        doc,
        "Grupo 1",
        "Autenticación y sistema general",
        "Login, pantalla de inicio, menú principal, usuarios, roles y acceso a los módulos.",
        "Cómo entra un usuario, qué ve al iniciar, qué roles existen y qué puede hacer cada rol.",
    )
    add_group_card(
        doc,
        "Grupo 2",
        "Laboratorio",
        "Carga, listado y gestión de turnos de laboratorio.",
        "Qué datos tiene un turno, quién lo carga, quién lo confirma, cómo se cancela y cómo se muestra el listado.",
    )
    add_group_card(
        doc,
        "Grupo 3",
        "Comedor",
        "Carga, listado y gestión de reservas o turnos del comedor.",
        "Qué datos necesita comedor, quién carga una reserva, cómo se confirma y cómo se controla la cantidad de alumnos.",
    )
    add_group_card(
        doc,
        "Grupo 4",
        "Notebooks y accesorios",
        "Registro y consulta de notebooks, accesorios, estados y movimientos de gabinete.",
        "Cómo se identifica cada equipo, qué estados puede tener, quién puede modificarlo y dónde está ubicado.",
    )

    doc.add_page_break()
    add_heading(doc, "Acuerdos Comunes", level=1)
    add_bullets(
        doc,
        [
            "Usar nombres claros en español.",
            "Pensar el sistema como una sola aplicación.",
            "Mantener pantallas parecidas entre los módulos.",
            "Usar los mismos nombres para botones comunes.",
            "Definir qué roles pueden ver o usar cada parte.",
            "Avisar si un módulo necesita información de otro.",
        ],
    )
    add_code_block(doc, ["Nuevo", "Guardar", "Cancelar", "Editar", "Eliminar", "Ver detalle", "Volver"])

    add_heading(doc, "Actividad", level=1)
    doc.add_paragraph("Cada grupo debe completar la siguiente ficha.")
    add_code_block(
        doc,
        [
            "Nombre del módulo:",
            "",
            "Integrantes:",
            "",
            "Objetivo del módulo:",
            "",
            "Problema que resuelve:",
            "",
            "Pantallas necesarias:",
            "",
            "Datos que se cargan:",
            "",
            "Datos que se muestran:",
            "",
            "Acciones del usuario:",
            "",
            "Estados posibles:",
            "",
            "Roles que usan este módulo:",
            "",
            "Información que necesita de otros módulos:",
            "",
            "Dudas o decisiones pendientes:",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "Para La Puesta En Común", level=1)
    doc.add_paragraph("Cada grupo deberá explicar brevemente:")
    add_bullets(
        doc,
        [
            "Qué módulo le tocó.",
            "Qué pantallas pensó.",
            "Qué datos necesita cargar.",
            "Qué acciones puede hacer el usuario.",
            "Qué roles usarían ese módulo.",
            "Qué dudas aparecieron.",
        ],
    )

    add_heading(doc, "Entrega Para La Próxima Clase", level=1)
    add_bullets(
        doc,
        [
            "Ficha del módulo completa.",
            "Boceto simple de las pantallas principales.",
            "Lista de dudas o decisiones para consultar.",
        ],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
