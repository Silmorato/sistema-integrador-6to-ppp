from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Clase de hoy - Organizacion del Proyecto Integrador.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(item)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(item)


def add_code_block(doc, lines):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        if index < len(lines) - 1:
            run.add_break()


def add_group_section(doc, title, purpose, must_think, data, states, questions, page_break=False):
    if page_break:
        doc.add_page_break()
    add_heading(doc, title, level=2)
    doc.add_paragraph(purpose)
    add_heading(doc, "Debe pensar", level=3)
    add_bullets(doc, must_think)
    add_heading(doc, "Datos posibles", level=3)
    add_code_block(doc, data)
    if states:
        add_heading(doc, "Estados posibles", level=3)
        add_code_block(doc, states)
    add_heading(doc, "Preguntas guia", level=3)
    add_code_block(doc, questions)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, before, after in [
        ("Heading 1", 20, 20, 6),
        ("Heading 2", 16, 18, 6),
        ("Heading 3", 14, 16, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Clase de hoy: Organización del Proyecto Integrador")
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = doc.add_paragraph()
    subtitle.add_run("Guía docente para presentar la organización del sistema y definir el trabajo por grupos.")

    add_heading(doc, "Objetivo De La Clase", level=1)
    doc.add_paragraph(
        "Que los grupos comprendan cómo se va a organizar el desarrollo del sistema integrador y definan con claridad qué debe hacer cada módulo."
    )
    add_code_block(
        doc,
        [
            "No vamos a construir cuatro sistemas separados.",
            "Vamos a construir un solo sistema integrador dividido en módulos.",
        ],
    )

    add_heading(doc, "Situación Del Proyecto", level=1)
    doc.add_paragraph(
        "El sistema integrador va a tener una pantalla de acceso, una pantalla de inicio y diferentes módulos para gestionar movimientos de la escuela."
    )

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    set_cell_text(hdr[0], "Módulo", bold=True)
    set_cell_text(hdr[1], "Responsabilidad principal", bold=True)
    set_cell_shading(hdr[0], "F2F4F7")
    set_cell_shading(hdr[1], "F2F4F7")
    rows = [
        ("Autenticación y sistema general", "Login, inicio, menú, usuarios, roles y acceso a módulos."),
        ("Laboratorio", "Carga, listado y gestión de turnos de laboratorio."),
        ("Comedor", "Carga, listado y gestión de reservas o turnos de comedor."),
        ("Notebooks y accesorios", "Gestion de equipos, accesorios, estados y movimientos."),
    ]
    for left, right in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], left)
        set_cell_text(cells[1], right)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    doc.add_paragraph(
        "En esta primera etapa vamos a trabajar sobre el frontend. El objetivo es lograr un prototipo navegable, con pantallas, formularios, listados y datos simulados."
    )

    add_heading(doc, "Explicación Inicial", level=1)
    doc.add_paragraph("Hasta ahora estuvimos trabajando en el análisis del sistema:")
    add_bullets(
        doc,
        [
            "Requerimientos funcionales y no funcionales.",
            "Historias de usuario y criterios de aceptación.",
            "Diagramas de flujo e ideas generales de funcionamiento.",
        ],
    )
    doc.add_paragraph(
        "Ahora empezamos una nueva etapa: pensar cómo se construye el sistema. Para trabajar entre varios grupos necesitamos una estructura común."
    )

    add_heading(doc, "Organización De Los Grupos", level=1)
    add_group_section(
        doc,
        "Grupo 1: Autenticación Y Sistema General",
        "Este grupo se encarga de la entrada al sistema y de la navegación general.",
        [
            "Pantalla de login.",
            "Pantalla de inicio.",
            "Menú principal y accesos a módulos.",
            "Roles de usuario y permisos visuales.",
            "Feed de novedades o calendario, si el tiempo alcanza.",
        ],
        ["Nombre", "Apellido", "Usuario", "Rol", "Estado"],
        [],
        [
            "Como ingresa un usuario al sistema",
            "Qué aparece en la pantalla de inicio",
            "Qué roles existen",
            "Qué puede ver o hacer cada rol",
        ],
    )
    add_group_section(
        doc,
        "Grupo 2: Laboratorio",
        "Este grupo se encarga de la gestión de turnos de laboratorio.",
        [
            "Formulario para cargar un turno.",
            "Listado de turnos.",
            "Edicion o cancelacion de turnos.",
            "Estados del turno y filtros posibles.",
        ],
        ["Fecha", "Hora", "Profesor", "Curso", "Laboratorio", "Motivo", "Estado", "Observaciones"],
        ["Pendiente", "Confirmado", "Cancelado", "Finalizado"],
        [
            "Quien puede pedir un turno",
            "Quien puede confirmarlo",
            "Qué datos son obligatorios",
            "Qué pasa si se cancela un turno",
        ],
        page_break=True,
    )
    add_group_section(
        doc,
        "Grupo 3: Comedor",
        "Este grupo se encarga de la gestión de reservas o turnos del comedor.",
        [
            "Formulario para cargar una reserva.",
            "Listado de reservas.",
            "Edicion o cancelacion de reservas.",
            "Estados de la reserva y filtros posibles.",
        ],
        ["Fecha", "Turno", "Curso", "Cantidad de alumnos", "Responsable", "Observaciones", "Estado"],
        ["Pendiente", "Confirmado", "Cancelado"],
        [
            "Quien puede cargar una reserva",
            "Qué información necesita comedor",
            "Como se controla la cantidad de alumnos",
            "Qué pasa si una reserva se cancela",
        ],
        page_break=True,
    )
    add_group_section(
        doc,
        "Grupo 4: Notebooks Y Accesorios",
        "Este grupo se encarga de registrar y consultar el estado de notebooks, accesorios y movimientos de gabinete.",
        [
            "Formulario para registrar notebook o accesorio.",
            "Listado de equipos.",
            "Cambio de estado.",
            "Asignacion a responsable o ubicacion.",
            "Historial simple de movimientos, si el tiempo alcanza.",
        ],
        ["Código interno", "Tipo", "Marca", "Modelo", "Estado", "Ubicación", "Responsable", "Observaciones"],
        ["Disponible", "En uso", "En reparación", "Fuera de servicio"],
        [
            "Como se identifica cada notebook",
            "Qué estados puede tener",
            "Quien puede cambiar el estado",
            "A quien puede estar asignada",
        ],
        page_break=True,
    )

    doc.add_page_break()
    add_heading(doc, "Acuerdos Comunes", level=1)
    add_bullets(
        doc,
        [
            "Usar nombres claros en español.",
            "Mantener una estructura similar de pantallas.",
            "Usar los mismos nombres para botones comunes.",
            "Definir datos, acciones, estados y roles.",
            "Avisar si un módulo necesita información de otro.",
        ],
    )
    add_code_block(doc, ["Nuevo", "Guardar", "Cancelar", "Editar", "Eliminar", "Ver detalle", "Volver"])

    add_heading(doc, "Actividad Para La Clase", level=1)
    doc.add_paragraph("Cada grupo debe completar una ficha de su módulo.")
    add_code_block(
        doc,
        [
            "Nombre del módulo:",
            "Integrantes:",
            "Objetivo del módulo:",
            "Problema que resuelve:",
            "Pantallas necesarias:",
            "Datos que se cargan:",
            "Datos que se muestran:",
            "Acciones del usuario:",
            "Estados posibles:",
            "Roles que usan este módulo:",
            "Información que necesita de otros módulos:",
            "Dudas o decisiones pendientes:",
        ],
    )

    add_heading(doc, "Puesta En Comun", level=1)
    doc.add_paragraph("Al final de la clase, cada grupo comparte:")
    add_bullets(
        doc,
        [
            "Qué módulo tiene.",
            "Qué pantallas pensó.",
            "Qué datos necesita cargar.",
            "Qué roles van a usarlo.",
            "Qué dudas aparecieron.",
        ],
    )

    add_heading(doc, "Cierre", level=1)
    add_code_block(
        doc,
        [
            "Hoy no programamos todavia.",
            "Hoy organizamos el proyecto para poder programar mejor.",
            "Un sistema no empieza solamente escribiendo código:",
            "empieza definiendo responsabilidades, pantallas, datos y acuerdos comunes.",
        ],
    )

    add_heading(doc, "Tarea Para La Proxima Clase", level=1)
    add_bullets(
        doc,
        [
            "Ficha del módulo completa.",
            "Boceto simple de las pantallas principales.",
            "Lista de dudas o decisiones que necesita consultar.",
        ],
    )
    doc.add_paragraph(
        "Para la clase de laboratorio del jueves se puede comenzar con Git, GitHub, ramas y versionado."
    )

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
